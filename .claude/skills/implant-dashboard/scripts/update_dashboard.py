#!/usr/bin/env python3
"""Merge monthly surgery-system exports into the implant dashboard and rebuild the page.

Usage:
  python3 update_dashboard.py --current <published dashboard .html or data .json> \
      --export <export.xlsx> [--export <export.xlsx> ...] --out <dir>

Writes <out>/data.json and <out>/implant-dashboard.html, and prints a JSON report
(per-month totals plus a check against the export's 실적요약 sheet).
Months present in an export replace the same months already in the dashboard, so
re-running with a later (closed) export of the same month is safe.
"""
import argparse, collections, datetime, json, os, re, sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, '..', 'assets', 'template.html')

MAKER = {'NEO': 'Neo', '유한양행': '유한_EVER', '신흥': '신흥_EVER', 'ASTRA': 'ASTRA',
         'DENTIUM': 'Den', '덴티움': 'Den', 'MEGAGEN': '메가젠', '메가젠': '메가젠'}
OSSTEM_LINE = {'TS3BA': 'TS3', 'TS3': 'TS3', 'TS4': 'TS4', 'TS5': 'TS5'}
SKIP_NAMES = {'E', 'x', 'X', '', 'fx', 'EXCL'}


def load_current(path):
    text = open(path, encoding='utf-8').read()
    if path.endswith('.json'):
        return json.loads(text)
    m = re.search(r'const RAW = (\{.*?\});\n', text, re.S)
    if not m:
        sys.exit(f'No "const RAW = {{...}};" block found in {path}')
    return json.loads(m.group(1))


def num(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return None


def vendor(maker, brand):
    maker = (maker or '').strip().upper() if maker and maker.isascii() else (maker or '').strip()
    if maker == 'OSSTEM':
        return 'Osstem', OSSTEM_LINE.get((brand or '').strip().upper(), '미기재')
    if not maker or maker == '확인안됨':
        return None, None
    return MAKER.get(maker, maker), None


def read_export(path, alias):
    wb = openpyxl.load_workbook(path, data_only=True)
    if 'View' not in wb.sheetnames:
        sys.exit(f'{path}: no "View" sheet. Is this a surgery-system export?')
    rows = list(wb['View'].iter_rows(values_only=True))
    head = rows[0]
    recs = []
    for r in rows[1:]:
        r = dict(zip(head, r))
        if not r.get('수술일자') or r.get('삭제상태') not in (None, '정상'):
            continue
        day = str(r['수술일자'])[:10]
        co, line = vendor(r.get('제조사'), r.get('브랜드'))
        yn = lambda k, who: (who or 'Y') if r.get(k) == 'Y' else 'x'
        recs.append({
            'month': day[:7], 'day': int(day[8:10]), 'chart': r.get('차트번호'),
            'kind': r.get('구분'), 'proc': 'free' if r.get('구분') == 'Free' else r.get('시술내용'),
            'dr': alias.get(r.get('담당의사'), r.get('담당의사')),
            'co': co, 'line': line,
            'fx': (str(r['Fixture']).upper() if r.get('Fixture') not in (None, '확인안됨') else None),
            'dia': num(r.get('직경')), 'len': num(r.get('길이')), 'qty': int(num(r.get('사용수량')) or 1),
            'bone': r.get('골이식'), 'ext': r.get('Ext') == 'Y', 'sinus': bool(r.get('SINUS')),
            'main': r.get('Main Assistant'),
            'hc': 'EXCL' if r.get('해피콜') == '제외' else yn('해피콜', r.get('해피콜 담당')),
            'sim': yn('Simulation', r.get('Simulation 담당')), 'guide': yn('Guide', r.get('Guide 담당')),
            'reason': r.get('제거원인'), 'closed': r.get('월마감') == '마감',
        })
    summary = {}
    if '실적요약' in wb.sheetnames:
        for row in wb['실적요약'].iter_rows(values_only=True):
            if row and row[0] and isinstance(row[1], (int, float)):
                summary[row[0]] = row[1]
    return recs, summary


def name(x):
    x = (str(x).strip() if x is not None else '')
    return None if x in SKIP_NAMES or x.startswith('=') else x


def fmt(v):
    return '%g' % v


def aggregate(recs):
    """Return per-month aggregates in the dashboard's RAW key format."""
    out = collections.defaultdict(lambda: collections.defaultdict(dict))
    by_month = collections.defaultdict(list)
    for d in recs:
        by_month[d['month']].append(d)
    for m, rs in by_month.items():
        A = out[m]
        mc, size, typ, rem, rei, line, staff = (collections.Counter() for _ in range(7))
        new = jg = rg = 0
        cases = collections.OrderedDict()
        for d in rs:
            cases.setdefault((d['day'], d['chart']), []).append(d)
            n, p = d['qty'], d['proc']
            if p == '신규' and d['co']:
                new += n
                mc[f"{m}|||{d['co']}"] += n
                if d['fx']:
                    typ[f"{m}|||{d['co']}|||{d['fx']}"] += n
                if d['dia'] and d['len']:
                    ln = d['len'] / 10 if d['len'] > 20 else d['len']
                    size[f"{m}|||{d['co']}|||{fmt(d['dia'])}|||{fmt(ln)}"] += n
                if d['co'] == 'Osstem':
                    line[f"{m}|||{d['line']}"] += n
            elif p == '신규':
                new += n
            if p == '제거':
                jg += n
                rem[f"{m}|||{d['co'] or '미기재'}"] += n
            if p in ('재식립', '제거&재식립'):
                rg += n
                rei[f"{m}|||{d['co'] or '미기재'}"] += n
        cl = list(cases.values())
        anyf = lambda rows, k: any(name(r[k]) for r in rows)
        for c in cl:
            a = name(c[0]['main'])
            if a:
                staff[f'{m}|||{a}'] += 1
        A['siknew'] = new; A['jegeo'] = jg; A['reimg'] = rg
        A['monthly_co'] = dict(mc); A['size_mco'] = dict(size); A['type_m'] = dict(typ)
        A['rem_m'] = dict(rem); A['rei_m'] = dict(rei); A['osstem_line'] = dict(line); A['staff_m'] = dict(staff)
        A['ops'] = dict(
            rows=len(rs), cases=len(cl),
            fixtures=sum(d['qty'] for d in rs if d['co'] and d['proc'] in ('신규', '재식립', '제거&재식립', 'free')),
            ins=sum(d['qty'] for d in rs if d['proc'] == '신규' and d['kind'] == '보험'),
            bone_s=sum(1 for d in rs if d['bone'] == '단순'), bone_c=sum(1 for d in rs if d['bone'] == '복잡'),
            ext=sum(1 for d in rs if d['ext']), sinus=sum(1 for d in rs if d['sinus']),
            hc=sum(1 for c in cl if anyf(c, 'hc')), hc_excl=sum(1 for c in cl if all(r['hc'] == 'EXCL' for r in c)),
            sim=sum(1 for c in cl if anyf(c, 'sim')), guide=sum(1 for c in cl if anyf(c, 'guide')),
            dr=dict(collections.Counter(d['dr'] for d in rs if d['dr'])),
            reasons=dict(collections.Counter(d['reason'] for d in rs if d['reason'])),
            free=sum(d['qty'] for d in rs if d['proc'] == 'free'),
            closed=all(d['closed'] for d in rs), last_day=max(d['day'] for d in rs),
        )
    return out


MONTH_DICTS = ['monthly_co', 'size_mco', 'type_m', 'rem_m', 'rei_m', 'osstem_line', 'staff_m']


def merge(raw, agg):
    for m, A in agg.items():
        for k in MONTH_DICTS:
            raw.setdefault(k, {})
            raw[k] = {kk: v for kk, v in raw[k].items() if not kk.startswith(m + '|||')}
            raw[k].update(A[k])
        raw['siknew_monthly'][m] = A['siknew']
        raw['jegeo_monthly'][m] = A['jegeo']
        raw['reimg_monthly'][m] = A['reimg']
        raw.setdefault('ops', {})[m] = A['ops']
    months = set(raw['months']) | set(agg)
    raw['months'] = sorted(months)
    tot = collections.Counter()
    for k, v in raw['monthly_co'].items():
        tot[k.split('|||')[1]] += v
    raw['companies'] = [c for c, _ in tot.most_common()]
    raw['updated'] = datetime.date.today().isoformat()
    return raw


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--current', required=True)
    ap.add_argument('--export', action='append', required=True)
    ap.add_argument('--out', required=True)
    a = ap.parse_args()
    raw = load_current(a.current)
    alias = raw.setdefault('config', {}).setdefault('doctor_alias', {})
    recs, report = [], {}
    known = {d for o in raw.get('ops', {}).values() for d in o.get('dr', {})}
    for p in a.export:
        r, summary = read_export(p, alias)
        recs += r
        for m, A in aggregate(r).items():
            o = A['ops']
            report[m] = dict(file=os.path.basename(p), new=A['siknew'], removal=A['jegeo'], reimplant=A['reimg'],
                             cases=o['cases'], fixtures=o['fixtures'], closed=o['closed'], last_day=o['last_day'],
                             check_cases=summary.get('수술건수'), check_fixtures=summary.get('Fixture'),
                             new_doctors=sorted({d['dr'] for d in r if d['month'] == m and d['dr'] and d['dr'] not in known}))
    if not recs:
        sys.exit('No usable rows in the export(s).')
    raw = merge(raw, aggregate(recs))
    os.makedirs(a.out, exist_ok=True)
    data = json.dumps(raw, ensure_ascii=False, separators=(',', ':'))
    open(os.path.join(a.out, 'data.json'), 'w', encoding='utf-8').write(data)
    page = open(TEMPLATE, encoding='utf-8').read().replace('__DATA__', data)
    open(os.path.join(a.out, 'implant-dashboard.html'), 'w', encoding='utf-8').write(page)
    print(json.dumps(report, ensure_ascii=False, indent=1))


if __name__ == '__main__':
    main()
